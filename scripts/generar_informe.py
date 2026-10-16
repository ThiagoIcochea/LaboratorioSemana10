#!/usr/bin/env python3
"""
Genera informe/Informe_Laboratorio_Monitoreo_Cloud.docx a partir de los
resultados reales (resultados/*.csv|json) y las evidencias (evidencias/*.png).
"""
import csv
import json
import pathlib

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parents[1]
RES, EV = ROOT / "resultados", ROOT / "evidencias"
OUT = ROOT / "informe/Informe_Laboratorio_Monitoreo_Cloud.docx"
REPO = "https://github.com/ThiagoIcochea/LaboratorioSemana10"

F = {r["fase"]: r for r in csv.DictReader(open(RES / "resumen_fases.csv"))}
B, I, R = F["linea_base"], F["incidente"], F["recuperacion"]
TL = {r["evento"]: r for r in csv.DictReader(open(RES / "timeline.csv"))}
AL = json.load(open(RES / "alertas_transiciones.json"))
CONS = json.load(open(RES / "consultas_promql.json"))


def n(v, k=1.0, d=1):
    return f"{float(v) * k:,.{d}f}".replace(",", " ").replace(".", ",")


def hora(ev):
    return TL[ev]["iso_lima"][11:19]


def first(alerta, estado):
    return next(a for a in AL if a["alerta"] == alerta and a["estado"] == estado)


def cons(cid, fase):
    return next(c for c in CONS if c["id"] == cid and c["fase"] == fase)


def ms(v):
    return n(v, 1000, 1) + " ms"


def pct(v, d=2):
    return n(v, 100, d) + " %"


# ---------------------------------------------------------------------------
doc = Document()
st = doc.styles
st["Normal"].font.name = "Calibri"
st["Normal"].font.size = Pt(11)
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.15
for lvl, size, color in ((1, 15, "1F3864"), (2, 12.5, "2E5597"), (3, 11.5, "2E5597")):
    h = st[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor.from_string(color)
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    rf = h.element.rPr.rFonts
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Calibri")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True

sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.3)
sec.top_margin = sec.bottom_margin = Cm(2.2)
TEXT_W = 21.0 - 4.6


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)


def field(par, instr, placeholder=""):
    r = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r2 = par.add_run()
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr; r2._r.append(t)
    r3 = par.add_run()
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r3._r.append(s)
    par.add_run(placeholder)
    r4 = par.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r4._r.append(e)


def P(text="", bold=False, italic=False, align=None, size=None, after=None, style=None, color=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_rich(p, text, bold, italic, size, color)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_rich(p, text, bold=False, italic=False, size=None, color=None):
    """Admite **negrita** y `código` en línea."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
        else:
            r = p.add_run(part); r.bold = bold
        r.italic = italic or r.italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def J(text):
    return P(text, align="j")


def bullet(text, num=False):
    p = doc.add_paragraph(style="List Number" if num else "List Bullet")
    add_rich(p, text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    return p


def H(text, lvl=1):
    return doc.add_heading(text, level=lvl)


def code(text, size=8.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, "F3F5F8")
    c.width = Cm(TEXT_W)
    lines = text.strip("\n").split("\n")
    p = c.paragraphs[0]
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


TAB_N = [0]
FIG_N = [0]


def caption_tab(title):
    TAB_N[0] += 1
    p = P(f"Tabla {TAB_N[0]}. ", bold=True, after=3)
    add_rich(p, title, italic=True)
    p.paragraph_format.keep_with_next = True


def table(header, rows, widths, title=None, source=None, size=9):
    if title:
        caption_tab(title)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3864")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            add_rich(cells[i].paragraphs[0], str(v), size=size)
            if ri % 2:
                shade(cells[i], "F2F5FA")
    for row in t.rows:
        # repetir cabecera
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    if source:
        P(source, italic=True, size=8.5, after=8, color="555555")
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, title, interpret, width=TEXT_W, source=None):
    FIG_N[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width))
    c = P(f"Figura {FIG_N[0]}. ", bold=True, align="c", after=1, size=9.5)
    add_rich(c, title, italic=True, size=9.5)
    c.paragraph_format.keep_with_next = True
    src = source or "Fuente: elaboración propia con datos reales de Prometheus 3.5.0 exportados del laboratorio."
    P(src, italic=True, align="c", size=8.5, after=3, color="555555").paragraph_format.keep_with_next = True
    P("Interpretación: " + interpret, align="j", size=10, after=10)
    return FIG_N[0]


SRC_UI = ("Fuente: captura auténtica de la interfaz web de Prometheus 3.5.0 del laboratorio "
          "(eje temporal en UTC; hora de Lima = UTC − 5 h).")

# ===========================================================================
# PORTADA
# ===========================================================================
P("UNIVERSIDAD TECNOLÓGICA DEL PERÚ", bold=True, align="c", size=18, after=2, color="1F3864")
P("Facultad de Ingeniería de Sistemas y Electrónica", align="c", size=12, after=30)
P("Curso: Herramientas de Desarrollo Profesional TIC", bold=True, align="c", size=13, after=40)
P("INFORME DE LABORATORIO N.º ____", bold=True, align="c", size=13, after=14, color="2E5597")
P("Monitoreo en la nube y gestión de métricas en bases de datos de series temporales",
  bold=True, align="c", size=20, after=8, color="1F3864")
P("Análisis de un incidente de degradación de una dependencia con Prometheus, PromQL y Grafana",
  italic=True, align="c", size=12, after=50)

ct = doc.add_table(rows=0, cols=2)
ct.style = "Table Grid"
ct.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in [("Docente", "Dr. Eleodoro Orbegoso Carrera"),
             ("Estudiante", "Icochea Rodriguez, Thiago Paolo"),
             ("Otros integrantes", "______________________________________"),
             ("Sección", "____________"),
             ("Equipo", "____________"),
             ("Número de laboratorio", "____"),
             ("Fecha de realización", "16/10/2026"),
             ("Repositorio", REPO)]:
    cells = ct.add_row().cells
    cells[0].text, cells[1].text = "", ""
    r = cells[0].paragraphs[0].add_run(k); r.bold = True
    shade(cells[0], "E8EEF7")
    cells[1].paragraphs[0].add_run(v)
    cells[0].width, cells[1].width = Cm(5.2), Cm(9.8)
P("", after=40)
P("Lima – Perú, 2026", align="c", size=12)
doc.add_page_break()

# ===========================================================================
# ÍNDICE
# ===========================================================================
P("Contenido", bold=True, size=15, color="1F3864", after=8)
field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Actualice el índice (clic derecho → Actualizar campo).")
doc.add_page_break()

# pie de página con numeración
for s in doc.sections:
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Laboratorio de Monitoreo Cloud – UTP   |   Página ")
    r.font.size = Pt(8.5)
    field(fp, "PAGE", "1")
    s.different_first_page_header_footer = True

# ===========================================================================
# 1. RESUMEN
# ===========================================================================
H("1. Resumen ejecutivo")
J(f"Este laboratorio tuvo como objetivo comprender y aplicar el monitoreo en la nube mediante métricas almacenadas en "
  f"una base de datos de series temporales. Se implementó una aplicación web instrumentada, node_exporter, Prometheus 3.5.0 "
  f"como recolector y TSDB, reglas de alerta y un dashboard de Grafana provisionado, empaquetados con Docker Compose para "
  f"GitHub Codespaces. Con tráfico constante de 8 peticiones por segundo se registró una línea base, luego se degradó de forma "
  f"controlada la dependencia interna inventario-db durante cinco minutos y finalmente se verificó la recuperación. "
  f"Durante el incidente, la latencia p95 pasó de {ms(B['latencia_p95_s'])} a {n(I['latencia_p95_s'], 1, 2)} s y los errores "
  f"5xx de {pct(B['ratio_errores_5xx'])} a {pct(I['ratio_errores_5xx'], 1)}, mientras la CPU del host se mantuvo por debajo "
  f"del {n(I['cpu_host_max_1m'], 100, 0)} % y la memoria disponible sobre el {n(I['mem_disponible_min_ratio'], 100, 0)} %. "
  f"La correlación de estas señales descartó la presión de recursos y señaló a la dependencia como causa. Las alertas "
  f"dispararon unos 75 segundos después del inicio. Se recomienda aplicar timeouts y un cortacircuitos hacia la dependencia, "
  f"además de alertar sobre su latencia.")

# ===========================================================================
# 2. OBJETIVOS
# ===========================================================================
H("2. Objetivos")
H("2.1. Objetivo general", 2)
J("Explicar los sistemas de monitoreo en la nube y analizar la gestión de métricas en bases de datos de series temporales "
  "mediante un caso de fallo de infraestructura reproducido de forma controlada (logro de la sesión).")
H("2.2. Objetivos específicos", 2)
for t in [
    "Implementar una arquitectura de monitoreo reproducible (fuente de métricas → Prometheus/TSDB → PromQL → Grafana → alertas) y explicar la función de cada componente.",
    "Construir un inventario de métricas reales con su significado, unidad, etiquetas y comportamiento esperado.",
    "Registrar la línea base del servicio y compararla con el intervalo de incidente y de recuperación mediante consultas PromQL.",
    "Contrastar hipótesis (CPU, memoria, E/S, dependencia) correlacionando al menos tres señales temporales.",
    "Proponer acciones correctivas, preventivas y reglas de alerta con umbrales justificados por los datos.",
]:
    bullet(t)
table(["Resultado específico de la guía", "Evidencia en este informe"], [
    ["Explica la arquitectura de un sistema de monitoreo en la nube", "Sección 4, Figura 1 y Tabla 3"],
    ["Distingue métricas de infraestructura y su dimensión temporal", "Sección 6.1, Tabla 5 (inventario)"],
    ["Analiza un fallo con datos temporales", "Secciones 6.2–6.6 y 7 (consultas, figuras, hoja de análisis)"],
    ["Propone acciones de respuesta y prevención", "Sección 8 y reglas de alerta del Anexo B"],
], [7.6, 8.8], "Correspondencia entre los resultados de aprendizaje y las evidencias")

# ===========================================================================
# 3. FUNDAMENTO
# ===========================================================================
H("3. Fundamento conceptual")
conceptos = [
    ("Monitoreo", "Proceso de obtener y seguir señales cuantitativas del comportamiento de aplicaciones e infraestructura para "
     "detectar desviaciones y apoyar decisiones operativas. Beyer et al. (2016) proponen vigilar cuatro «señales doradas»: "
     "latencia, tráfico, errores y saturación; este laboratorio mide las cuatro."),
    ("Observabilidad", "Capacidad de explicar el estado interno de un sistema a partir de los datos que emite (métricas, "
     "registros y trazas). A diferencia del monitoreo de umbrales conocidos, permite investigar fallos no previstos "
     "formulando nuevas preguntas sobre la telemetría (Majors et al., 2022)."),
    ("Métrica", "Medición numérica observada en el tiempo. Prometheus define cuatro tipos: contador (solo crece, p. ej. "
     "peticiones totales), gauge (sube y baja, p. ej. peticiones en curso), histograma (distribución en cubetas, p. ej. "
     "latencias) y resumen (Prometheus Authors, 2026b)."),
    ("Serie temporal", "Secuencia de muestras (marca de tiempo, valor). En Prometheus cada serie se identifica de forma única "
     "por el nombre de la métrica y su conjunto de etiquetas clave/valor; por ejemplo, "
     "`app_http_requests_total{endpoint=\"/api/checkout\",status=\"503\"}` es una serie distinta de la de estado 200 "
     "(Prometheus Authors, 2026b)."),
    ("TSDB", "Base de datos optimizada para datos indexados por tiempo y consultas por rangos. Prometheus incluye una TSDB "
     "local que agrupa las muestras en bloques de dos horas, usa un registro de escritura anticipada (WAL) para tolerar "
     "caídas y aplica una política de retención (Prometheus Authors, 2026d). En el laboratorio la retención es de 2 días."),
    ("Prometheus", "Sistema de monitoreo de código abierto que recolecta métricas mediante el modelo pull: consulta "
     "periódicamente el endpoint HTTP /metrics de cada objetivo, almacena las muestras en su TSDB y evalúa reglas de "
     "registro y de alerta (Prometheus Authors, 2026a)."),
    ("PromQL", "Lenguaje funcional de consulta de Prometheus. Opera sobre vectores instantáneos y de rango; funciones como "
     "`rate()` convierten contadores en tasas por segundo, `histogram_quantile()` estima percentiles a partir de cubetas y "
     "los agregadores (`sum by`, `avg by`) resumen conservando dimensiones (Prometheus Authors, 2026c)."),
    ("Grafana", "Plataforma de visualización que se conecta a Prometheus como fuente de datos, ejecuta consultas PromQL y "
     "presenta los resultados en paneles (Grafana Labs, 2026a)."),
    ("Dashboard", "Conjunto organizado de paneles que permite observar tendencias, picos y relaciones entre señales en una "
     "misma ventana de tiempo. En este laboratorio se provisiona como código (JSON), lo que lo hace reproducible "
     "(Grafana Labs, 2026b)."),
    ("Alertamiento", "Reglas que evalúan una condición PromQL de forma periódica; si se cumple durante el tiempo indicado en "
     "`for`, la alerta pasa de pending a firing (Prometheus Authors, 2026e). Una alerta indica que algo requiere atención; "
     "no sustituye el análisis causal."),
]
table(["Concepto", "Explicación"], [[f"**{a}**", b] for a, b in conceptos], [3.2, 13.2],
      "Conceptos fundamentales del laboratorio", "Fuente: elaboración propia a partir de las referencias citadas.", size=9.5)

# ===========================================================================
# 5 (4). ENTORNO Y ARQUITECTURA
# ===========================================================================
H("4. Entorno y arquitectura")
H("4.1. Entorno de ejecución", 2)
J("El proyecto se diseñó para **GitHub Codespaces**, el entorno de desarrollo en la nube de GitHub que ejecuta un contenedor "
  "de desarrollo definido en `.devcontainer/devcontainer.json` (GitHub, 2026). La configuración usa la imagen oficial de Python "
  "con la característica *Docker-in-Docker*, solicita 2 núcleos y 4 GB (dentro de la cuota gratuita) y reenvía los puertos 8000, "
  "9090 y 3000. Con **Docker Compose** (Docker Inc., 2026) se levantan cuatro servicios: app, node-exporter, prometheus y grafana.")
J("**Entorno donde se obtuvieron los resultados.** La ejecución real se realizó en una máquina virtual Linux en la nube "
  "(2 vCPU, 7,8 GiB de RAM) asociada a la sesión de trabajo. Ese entorno no podía descargar imágenes de Docker Hub ni el "
  "instalador de Grafana (la red las bloqueaba), por lo que se ejecutaron como procesos los **binarios oficiales** de "
  "Prometheus 3.5.0 y node_exporter 1.9.1 publicados en GitHub, con la misma configuración del repositorio "
  "(`scripts/run_local_sin_docker.sh`). Grafana no pudo ejecutarse allí: su datasource y su dashboard quedaron provisionados "
  "como código y se ejecutan con `docker compose up` en Codespaces. Por eso las figuras de este informe provienen de la interfaz "
  "real de Prometheus y de gráficos construidos con los datos exportados de su API; no se presentan capturas de Grafana "
  "simuladas. No se usaron servicios de pago.")
table(["Componente", "Nombre lógico / dirección", "Versión", "Función en el flujo"], [
    ["App «catalogo»", "app:8000 (localhost:8000)", "Python 3, Flask 3.1.2, prometheus-client 0.23.1, Gunicorn 23.0.0", "Fuente de métricas de aplicación (/metrics)"],
    ["Dependencia inventario-db", "interna a la app", "simulada", "Origen del fallo controlado"],
    ["node_exporter", "node-exporter:9100", "1.9.1", "Exportador de métricas del host"],
    ["Prometheus", "prometheus:9090", "3.5.0", "Recolección (pull, 5 s), TSDB, PromQL, reglas"],
    ["Grafana", "grafana:3000", "12.1.1 (provisionado)", "Visualización (dashboard de 12 paneles)"],
    ["Generador de tráfico", "scripts/loadgen.py", "—", "8 req/s: 70 % productos, 20 % checkout, 10 % raíz"],
], [3.2, 3.6, 4.6, 5.0], "Componentes, nombres lógicos y versiones (sin credenciales)",
      "Fuente: elaboración propia. Los nombres lógicos corresponden a Docker Compose; entre paréntesis, la dirección usada en la ejecución.")

H("4.2. Arquitectura implementada y flujo de métricas", 2)
figure(EV / "arquitectura/arquitectura.png", "Arquitectura del laboratorio y flujo de métricas",
       "el generador produce tráfico HTTP hacia la app; cada petición consulta a inventario-db. La app y node_exporter exponen "
       "/metrics, que Prometheus recolecta cada 5 s (modelo pull) y guarda en su TSDB. El motor PromQL alimenta las reglas de "
       "registro y de alerta, y Grafana consulta la API de Prometheus. Las alertas y paneles sirven de entrada al diagnóstico.",
       width=11.5, source="Fuente: elaboración propia (Graphviz, evidencias/arquitectura/arquitectura.dot).")
J("La aplicación expone métricas siguiendo las recomendaciones de instrumentación: un contador de peticiones con etiquetas "
  "`method`, `endpoint` y `status`; un histograma de latencia por endpoint; un histograma y un contador de errores de la "
  "dependencia; y un gauge de concurrencia. El modo de fallo se activa con `POST /admin/fault?enable=1` y solo afecta a "
  "inventario-db: añade una latencia de 0,6 a 1,4 s (con `time.sleep`, que **no consume CPU**) y un 30 % de errores por "
  "timeout. La ruta raíz `/` no consulta la dependencia, lo que sirve como grupo de control.")

H("4.3. Configuraciones esenciales", 2)
code("""# prometheus/prometheus.yml (extracto)
global:
  scrape_interval: 5s        # resolución para ver incidentes de pocos minutos
  evaluation_interval: 5s
rule_files: [/etc/prometheus/alerts.yml]
scrape_configs:
  - job_name: catalogo-app
    static_configs: [{targets: ["app:8000"], labels: {service: catalogo}}]
  - job_name: node
    static_configs: [{targets: ["node-exporter:9100"]}]
  - job_name: prometheus
    static_configs: [{targets: ["prometheus:9090"]}]""")
J("Se eligió un intervalo de 5 s porque el incidente dura cinco minutos: genera 60 muestras por serie durante el fallo sin "
  "sobrecargar el entorno (el proceso de la app consumió ≈ 0,012 núcleos). Los almacenes de Prometheus y Grafana usan "
  "volúmenes de Docker y la contraseña de Grafana se toma de una variable de entorno, no del repositorio.")

# ===========================================================================
# 6 (5). PROCEDIMIENTO
# ===========================================================================
H("5. Procedimiento")
pasos = [
    ("Verificación de herramientas", "Se comprobó la disponibilidad de Docker 29.8.2, Docker Compose 5.5.1, Python 3 y "
     "Chromium sin interfaz. La descarga de imágenes desde Docker Hub y de Grafana fue rechazada por la red (HTTP 403), por "
     "lo que se usó el modo sin Docker descrito en 4.1."),
    ("Construcción del proyecto", "Se crearon la app instrumentada, `prometheus.yml`, `alerts.yml`, el dashboard y el "
     "datasource de Grafana, `docker-compose.yml`, el devcontainer y los scripts de cada fase."),
    ("Validación de configuración", "`promtool check config` y `promtool check rules` devolvieron SUCCESS (10 reglas)."),
    ("Arranque y validación", "Se ejecutó `scripts/run_local_sin_docker.sh` y `scripts/validar_entorno.sh`: los tres "
     "objetivos quedaron en estado UP, la última muestra tenía 4,9 s de antigüedad y las reglas en estado inactive/ok."),
    ("Calentamiento", f"60 s de tráfico ({hora('calentamiento_inicio')}) para llenar las ventanas de `rate()`."),
    ("Línea base", f"`ejecutar_baseline.sh 360`: tráfico normal de {hora('baseline_inicio')} a {hora('baseline_fin')}."),
    ("Incidente", f"`ejecutar_incidente.sh 300`: activa el fallo a las {hora('incidente_inicio')}, mantiene el mismo tráfico y "
     f"lo desactiva a las {hora('incidente_fin')}."),
    ("Recuperación", f"`ejecutar_recuperacion.sh 360`: tráfico normal de {hora('recuperacion_inicio')} a {hora('recuperacion_fin')}."),
    ("Exportación y evidencias", "`exportar_metricas.sh` (series cada 5 s, 57 consultas PromQL por fase, transiciones de "
     "alertas), `graficar.py` (gráficos) y `capturas_prometheus.py` (capturas de la interfaz real)."),
]
table(["N.º", "Paso", "Acción y verificación"], [[str(i + 1), a, b] for i, (a, b) in enumerate(pasos)], [1.0, 3.6, 11.8],
      "Secuencia reproducible ejecutada (hora de Lima, 16/10/2026)")
P("Comandos para reproducir en GitHub Codespaces:", bold=True, after=3)
code("""docker compose up -d --build
bash scripts/validar_entorno.sh
bash scripts/ejecutar_experimento.sh      # base 360 s -> incidente 300 s -> recuperación 360 s -> exportación
python3 scripts/graficar.py && python3 scripts/capturas_prometheus.py
python3 scripts/generar_informe.py
docker compose down -v                    # y detener el Codespace""")
figure(EV / "arquitectura/prometheus_targets.png", "Objetivos de recolección de Prometheus en estado UP",
       "los tres objetivos (catalogo-app, node y prometheus) se recolectan correctamente y sin errores, lo que confirma que "
       "existen datos recientes antes de iniciar las mediciones.", width=14.5, source=SRC_UI)

# ===========================================================================
# 7 (6). RESULTADOS
# ===========================================================================
H("6. Resultados")
H("6.1. Inventario de métricas", 2)
table(["Métrica", "Qué representa", "Unidad / tipo", "Etiquetas relevantes", "Comportamiento esperado"], [
    ["`up`", "Si Prometheus alcanzó el objetivo en el último scrape", "0/1 (gauge sintético)", "job, instance", "Constante en 1"],
    ["`app_http_requests_total`", "Peticiones HTTP atendidas", "peticiones (contador)", "method, endpoint, status", "Crece ≈ 8/s; 5xx casi nulos"],
    ["`app_http_request_duration_seconds`", "Distribución de la latencia de respuesta", "segundos (histograma)", "endpoint, le", "p95 < 100 ms"],
    ["`app_dependency_duration_seconds`", "Latencia de llamadas a inventario-db", "segundos (histograma)", "dependency, le", "p95 ≈ 35–50 ms"],
    ["`app_dependency_errors_total`", "Errores de la dependencia", "errores (contador)", "dependency, reason", "≈ 0,2 % transitorios"],
    ["`app_http_inflight_requests`", "Peticiones en curso", "peticiones (gauge)", "—", "0–2"],
    ["`node_cpu_seconds_total`", "Tiempo de CPU por modo", "segundos (contador)", "cpu, mode", "idle ≈ 95 %"],
    ["`node_memory_MemAvailable_bytes`", "Memoria disponible del host", "bytes (gauge)", "instance", "> 85 % del total"],
    ["`node_disk_io_time_seconds_total`", "Tiempo de disco ocupado", "segundos (contador)", "device", "Cercano a 0"],
    ["`process_cpu_seconds_total`", "CPU consumida por el proceso de la app", "segundos (contador)", "job", "≈ 0,01 núcleos"],
    ["`process_resident_memory_bytes`", "Memoria residente de la app", "bytes (gauge)", "job", "≈ 30–35 MB estable"],
    ["`app_fault_injection_active`", "Marcador del experimento (no se usa como evidencia diagnóstica)", "0/1 (gauge)", "—", "0 salvo en el incidente"],
], [4.3, 3.9, 2.6, 2.6, 3.0], "Inventario de métricas reales expuestas", size=8.5,
      source="Fuente: elaboración propia a partir del endpoint /metrics de la app y de node_exporter.")
J("Para el incidente se seleccionaron cinco métricas relacionadas: latencia de la dependencia, latencia HTTP, tasa de errores, "
  "CPU del host y memoria disponible. Las dos últimas sirven para contrastar la hipótesis de presión de recursos sugerida por la guía.")

H("6.2. Línea base", 2)
J(f"Intervalo: **{B['inicio_lima']} a {B['fin_lima'][11:]}** (hora de Lima; {B['duracion_s']} s). Con tráfico constante, "
  f"el servicio atendió ≈ {n(B['peticiones_totales'], 1, 0)} peticiones ({n(B['rps_promedio'], 1, 2)} req/s), con latencia media de "
  f"{ms(B['latencia_media_s'])}, p95 de {ms(B['latencia_p95_s'])} y {pct(B['ratio_errores_5xx'], 3)} de respuestas 5xx (un error "
  f"transitorio de la dependencia). La CPU del host promedió {pct(B['cpu_host_media'], 1)}, la memoria disponible no bajó del "
  f"{pct(B['mem_disponible_min_ratio'], 1)} y la concurrencia máxima fue {n(B['inflight_max'], 1, 0)}. Este es el patrón de referencia.")
figure(EV / "baseline/fig_up.png", "Disponibilidad de los objetivos durante todo el experimento (up)",
       "up permaneció en 1 para la app y para node_exporter en las tres fases. El objetivo nunca dejó de responder a Prometheus, "
       "por lo que el incidente no fue una caída del servicio, sino una degradación.")
figure(EV / "baseline/prom_p95_baseline.png", "Latencia p95 (ventana de 1 min) durante la línea base en la interfaz de Prometheus",
       f"el p95 oscila entre unos 50 y 70 ms sin tendencia, consistente con la latencia normal de la dependencia "
       f"(p95 de {ms(B['dependencia_p95_s'])}).", width=14.5, source=SRC_UI)

H("6.3. Incidente", 2)
J(f"**Cambio introducido:** a las {hora('incidente_inicio')} se activó el modo de fallo de inventario-db (latencia de 0,6–1,4 s y "
  f"30 % de timeouts), sin modificar el tráfico ni los recursos. **Duración:** {I['duracion_s']} s. **Recuperación:** a las "
  f"{hora('incidente_fin')} se desactivó el fallo. El generador de tráfico registró, desde el lado del cliente, 2 400 peticiones, "
  f"778 errores (32,4 %) y latencia media de 1,063 s, coherente con lo medido por Prometheus.")
table(["Indicador (consulta)", "Línea base", "Incidente", "Recuperación"], [
    ["Peticiones por segundo (Q2b)", n(B["rps_promedio"], 1, 2), n(I["rps_promedio"], 1, 2), n(R["rps_promedio"], 1, 2)],
    ["Errores 5xx (Q3)", pct(B["ratio_errores_5xx"], 3), pct(I["ratio_errores_5xx"]), pct(R["ratio_errores_5xx"], 3)],
    ["Latencia media HTTP (Q4a)", ms(B["latencia_media_s"]), ms(I["latencia_media_s"]), ms(R["latencia_media_s"])],
    ["Latencia p95 HTTP (Q4b)", ms(B["latencia_p95_s"]), ms(I["latencia_p95_s"]), ms(R["latencia_p95_s"])],
    ["Máximo del p95 de 1 min (Q5)", ms(B["latencia_p95_max_1m_s"]), ms(I["latencia_p95_max_1m_s"]), ms(R["latencia_p95_max_1m_s"]) + " *"],
    ["p95 de inventario-db (Q6)", ms(B["dependencia_p95_s"]), ms(I["dependencia_p95_s"]), ms(R["dependencia_p95_s"])],
    ["Errores de inventario-db (Q7)", "≈ " + n(B["dependencia_errores"], 1, 0), "≈ " + n(I["dependencia_errores"], 1, 0), "≈ " + n(R["dependencia_errores"], 1, 0)],
    ["CPU host media / máx. 1 min (Q8a/b)", f"{pct(B['cpu_host_media'], 1)} / {pct(B['cpu_host_max_1m'], 1)}",
     f"{pct(I['cpu_host_media'], 1)} / {pct(I['cpu_host_max_1m'], 1)}", f"{pct(R['cpu_host_media'], 1)} / {pct(R['cpu_host_max_1m'], 1)}"],
    ["CPU del proceso app (Q8c)", n(B["cpu_proceso_app_nucleos"], 1, 4) + " núcleos", n(I["cpu_proceso_app_nucleos"], 1, 4) + " núcleos", n(R["cpu_proceso_app_nucleos"], 1, 4) + " núcleos"],
    ["Memoria disponible mínima (Q9a)", pct(B["mem_disponible_min_ratio"], 1), pct(I["mem_disponible_min_ratio"], 1), pct(R["mem_disponible_min_ratio"], 1)],
    ["RSS máx. de la app (Q9b)", n(B["rss_app_max_mb"], 1, 1) + " MB", n(I["rss_app_max_mb"], 1, 1) + " MB", n(R["rss_app_max_mb"], 1, 1) + " MB"],
    ["Concurrencia máxima (Q10)", n(B["inflight_max"], 1, 0), n(I["inflight_max"], 1, 0), n(R["inflight_max"], 1, 0)],
    ["Utilización máx. de E/S de disco (Q11)", pct(B["disk_io_util_max"]), pct(I["disk_io_util_max"]), pct(R["disk_io_util_max"])],
    ["Disponibilidad mínima up (Q1)", n(B["disponibilidad_min_up"], 1, 0), n(I["disponibilidad_min_up"], 1, 0), n(R["disponibilidad_min_up"], 1, 0)],
], [6.0, 3.4, 3.4, 3.6], "Comparación de indicadores por fase (consultas evaluadas sobre toda la ventana de cada fase)", size=8.5,
      source="Fuente: resultados/resumen_fases.csv. (*) El máximo de la recuperación incluye el primer minuto, cuya ventana "
             "rate[1m] todavía contiene muestras del incidente. Los conteos con «≈» provienen de increase(), que extrapola.")
figure(EV / "consultas/fig_rps.png", "Tasa de peticiones por código de estado",
       "el tráfico total se mantuvo en ≈ 8 req/s en las tres fases; durante el incidente una parte de las respuestas pasó de 2xx a "
       "5xx. Que la carga no cambiara descarta un pico de tráfico como causa.")
figure(EV / "consultas/fig_errores.png", "Proporción de respuestas 5xx con el umbral de alerta",
       f"la proporción subió desde ≈ 0 % hasta estabilizarse en 31–35 % (promedio de la fase {pct(I['ratio_errores_5xx'], 1)}) y "
       "volvió a la base en unos 60 s tras desactivar el fallo; la rampa de subida y bajada refleja la ventana de 1 minuto de rate().")
figure(EV / "consultas/fig_latencia.png", "Latencia HTTP: percentiles p50, p95 y p99",
       f"los tres percentiles saltaron en el mismo intervalo de 5 s en que se activó el fallo; el p95 se mantuvo cerca de "
       f"{n(I['latencia_p95_s'], 1, 1)} s. Como los percentiles se estiman por interpolación entre cubetas (1,5 y 2,5 s), su valor es "
       "aproximado, pero el cambio de magnitud (≈ 40 veces) es inequívoco.")
figure(EV / "consultas/fig_dependencia.png", "Latencia p95 y errores por segundo de inventario-db",
       f"el p95 de la dependencia pasó de {ms(B['dependencia_p95_s'])} a {ms(I['dependencia_p95_s'])} y aparecieron ≈ 2,6 errores/s "
       "de timeout. Es la señal más cercana al origen del fallo y la primera en superar su umbral.")
figure(EV / "consultas/fig_recursos.png", "CPU y memoria del host y del proceso de la aplicación",
       f"la CPU del host promedió {pct(I['cpu_host_media'], 1)} durante el incidente (menos que en la línea base) y el proceso de "
       f"la app consumió lo mismo en las tres fases (≈ {n(I['cpu_proceso_app_nucleos'], 1, 3)} núcleos). La memoria disponible y el "
       "RSS no cambiaron de forma relevante. El pico de CPU antes de las 18:31 corresponde al arranque de los servicios, fuera del experimento.")
figure(EV / "incidente/fig_inflight.png", "Peticiones en curso",
       f"la concurrencia subió de 1 a un máximo de {n(I['inflight_max'], 1, 0)}. Por la ley de Little, 8 req/s × 1,06 s ≈ 8,5 "
       "peticiones simultáneas: el aumento es consecuencia de respuestas más lentas, no de más tráfico.")
figure(EV / "incidente/prom_dependencia_p95.png", "p95 de inventario-db en la interfaz de Prometheus (regla de registro)",
       "la serie job:app_dependency_latency_seconds:p95_1m reproduce el escalón observado en los datos exportados, lo que confirma "
       "la trazabilidad entre la TSDB, las consultas y los gráficos.", width=14.5, source=SRC_UI)

H("6.4. Recuperación", 2)
J(f"Tras desactivar el fallo a las {hora('incidente_fin')}, la latencia de la dependencia volvió a su nivel normal en el siguiente "
  f"intervalo de recolección, y las señales basadas en ventanas de 1 minuto regresaron a la base en unos 60 s. En la fase de "
  f"recuperación completa el p95 HTTP fue {ms(R['latencia_p95_s'])} (base: {ms(B['latencia_p95_s'])}) y los errores "
  f"{pct(R['ratio_errores_5xx'], 2)}, valores equivalentes a la línea base. Los pocos errores restantes se explican por la tasa "
  f"transitoria normal de la dependencia y por peticiones iniciadas durante el fallo que terminaron después del corte.")
figure(EV / "incidente/prom_latencia_p95.png", "Ciclo completo de la latencia p95: línea base, incidente y recuperación",
       "la captura de la interfaz de Prometheus muestra el escalón de ≈ 0,05 s a ≈ 2,2 s y el retorno completo al nivel inicial, "
       "sin degradación residual.", width=14.5, source=SRC_UI)

H("6.5. Consultas PromQL", 2)
J("Cada consulta se evaluó con la API de Prometheus al final de cada fase, usando como ventana de rango la duración completa de la "
  "fase (Δ = 361 s, 302 s y 360 s). Los resultados completos están en `resultados/consultas_promql.json` (57 evaluaciones).")
q_rows = [
    ["Q1", 'min_over_time(up{job="catalogo-app"}[Δ])', "Disponibilidad del objetivo", "1 / 1 / 1",
     "El objetivo nunca dejó de ser recolectado."],
    ["Q2b", "sum(rate(app_http_requests_total[Δ]))", "Tasa por segundo (rate)",
     f"{n(B['rps_promedio'],1,2)} / {n(I['rps_promedio'],1,2)} / {n(R['rps_promedio'],1,2)}", "La carga fue constante."],
    ["Q2c", "sum by (endpoint, status) (rate(app_http_requests_total[Δ]))", "Tasa por endpoint y estado",
     "503 en incidente: productos 1,65/s; checkout 0,94/s; raíz 0", "Solo fallan los endpoints que usan inventario-db."],
    ["Q3", 'sum(increase(…{status=~"5.."}[Δ])) / sum(increase(…[Δ]))', "Tasa de errores",
     f"{pct(B['ratio_errores_5xx'],3)} / {pct(I['ratio_errores_5xx'])} / {pct(R['ratio_errores_5xx'],2)}", "Unas 900 veces más errores en el incidente."],
    ["Q4b", "histogram_quantile(0.95, sum by (le) (rate(…_bucket[Δ])))", "Latencia p95",
     f"{ms(B['latencia_p95_s'])} / {ms(I['latencia_p95_s'])} / {ms(R['latencia_p95_s'])}", "Degradación severa y reversible."],
    ["Q4c", "histogram_quantile(0.95, sum by (endpoint, le) (rate(…_bucket[Δ])))", "Agrupación by endpoint",
     "raíz 9,5 ms en las 3 fases; checkout 91 ms → 2 490 ms", "La ruta sin dependencia no se degradó; checkout (2 llamadas) es la más afectada."],
    ["Q5", "max_over_time(job:app_http_latency_seconds:p95_1m[Δ])", "Máximo temporal",
     f"{ms(B['latencia_p95_max_1m_s'])} / {ms(I['latencia_p95_max_1m_s'])}", "Pico del incidente ≈ 31 veces el pico base."],
    ["Q6", "histogram_quantile(0.95, sum by (le) (rate(app_dependency_duration_seconds_bucket[Δ])))", "Latencia de dependencia",
     f"{ms(B['dependencia_p95_s'])} / {ms(I['dependencia_p95_s'])} / {ms(R['dependencia_p95_s'])}", "Origen de la latencia observada."],
    ["Q8a", '1 - avg by (instance) (rate(node_cpu_seconds_total{mode="idle"}[Δ]))', "Uso de CPU (avg by)",
     f"{pct(B['cpu_host_media'],1)} / {pct(I['cpu_host_media'],1)} / {pct(R['cpu_host_media'],1)}", "Sin presión de CPU."],
    ["Q9a", "min_over_time((MemAvailable / MemTotal)[Δ:5s])", "Uso de memoria (subconsulta)",
     f"{pct(B['mem_disponible_min_ratio'],1)} / {pct(I['mem_disponible_min_ratio'],1)} / {pct(R['mem_disponible_min_ratio'],1)}", "Sin presión de memoria."],
]
table(["ID", "PromQL", "Objetivo", "Resultado (base / incidente / recuperación)", "Interpretación"], q_rows,
      [1.0, 5.2, 2.6, 3.9, 3.7], "Consultas PromQL principales, resultados e interpretación", size=8,
      source="Fuente: resultados/consultas_promql.json. «…» abrevia el nombre de métrica indicado en el Anexo A.")
figure(EV / "consultas/prom_max_over_time_p95.png", "Consulta max_over_time sobre el p95 en la interfaz de Prometheus",
       "max_over_time con ventana de 5 min conserva el pico durante 5 minutos después del incidente; es útil para revisar "
       "periodos pasados sin perder máximos breves, pero no debe usarse para detectar la recuperación en tiempo real.",
       width=14.5, source=SRC_UI)

H("6.6. Alertas", 2)
fe, fl, fd = first("TasaErroresAlta", "firing"), first("LatenciaP95Alta", "firing"), first("DependenciaLenta", "firing")
pe = first("LatenciaP95Alta", "pending")
table(["Alerta", "Condición (for)", "Pending desde", "Firing desde", "Firing hasta", "Resultado"], [
    ["DependenciaLenta", "p95 dependencia > 250 ms (1 m)", first("DependenciaLenta", "pending")["primero"][11:], fd["primero"][11:], fd["ultimo"][11:], "Disparó"],
    ["LatenciaP95Alta", "p95 HTTP > 500 ms (1 m)", pe["primero"][11:], fl["primero"][11:], fl["ultimo"][11:], "Disparó"],
    ["TasaErroresAlta", "5xx > 5 % (1 m)", first("TasaErroresAlta", "pending")["primero"][11:], fe["primero"][11:], fe["ultimo"][11:], "Disparó"],
    ["CPUHostAlta", "CPU host > 85 % (2 m)", "—", "—", "—", "No disparó (correcto)"],
    ["MemoriaHostBaja", "Mem. disponible < 10 % (2 m)", "—", "—", "—", "No disparó (correcto)"],
    ["ObjetivoCaido", "up == 0 (30 s)", "—", "—", "—", "No disparó (correcto)"],
], [3.0, 3.8, 2.2, 2.2, 2.2, 3.0], "Transiciones de las alertas durante el experimento (hora de Lima)", size=8.5,
      source="Fuente: resultados/alertas_transiciones.json (serie ALERTS de Prometheus, paso 5 s).")
J(f"Las alertas pasaron a *pending* 14–19 s después de activar el fallo y a *firing* unos 74–79 s después (por la cláusula "
  f"`for: 1m`). Se resolvieron entre 52 y 57 s después de desactivar el fallo, cuando las ventanas de 1 minuto dejaron de contener "
  f"muestras del incidente. Las alertas de recursos no se activaron, lo que coincide con los datos de CPU y memoria.")
figure(EV / "alertas/fig_alertas.png", "Número de alertas en estado firing",
       "las tres alertas de servicio estuvieron activas solo durante el incidente (más el desfase de las ventanas); fuera de él el "
       "conteo es 0, es decir, no hubo falsos positivos en la línea base ni en la recuperación.")
figure(EV / "alertas/prom_series_alerts.png", "Alertas por estado: sum by (alertstate) (ALERTS), gráfico apilado",
       "(azul: pending; verde: firing). Unos segundos después del inicio aparecen alertas en pending y, un minuto después, las tres pasan a firing; desaparecen "
       "cerca de un minuto después de retirar el fallo. Verifica que las reglas se evaluaron en el entorno real.", width=14.5, source=SRC_UI)

# ===========================================================================
# 8 (7). ANÁLISIS DEL FALLO
# ===========================================================================
H("7. Análisis del fallo")
H("7.1. Línea temporal", 2)
table(["Hora (Lima)", "Evento u observación", "Fuente"], [
    [hora("baseline_inicio"), "Inicio de la línea base; señales estables", "timeline.csv"],
    [hora("incidente_inicio"), "Se activa el fallo de inventario-db (cambio declarado)", "timeline.csv"],
    ["18:37:48–18:37:53", "p95 de la dependencia y p95 HTTP aumentan en el mismo intervalo de recolección", "Figuras 7, 8 y 16"],
    [pe["primero"][11:], "LatenciaP95Alta y DependenciaLenta en pending", "ALERTS"],
    ["≈ 18:38:50", "La proporción de 5xx se estabiliza en 31–35 % (fin de la rampa de 1 min)", "Figura 6"],
    [fd["primero"][11:], "DependenciaLenta y LatenciaP95Alta en firing", "ALERTS"],
    [fe["primero"][11:], "TasaErroresAlta en firing", "ALERTS"],
    [hora("incidente_fin"), "Se desactiva el fallo; la latencia de la dependencia se normaliza", "timeline.csv"],
    [f"{fe['ultimo'][11:]}–{fd['ultimo'][11:]}", "Se resuelven las tres alertas", "ALERTS"],
    [hora("recuperacion_fin"), "Fin de la recuperación; indicadores iguales a la línea base", "Tabla 6"],
], [3.2, 9.6, 3.6], "Línea temporal del incidente")

H("7.2. Hoja de análisis del incidente", 2)
table(["Elemento", "Registro del estudiante"], [
    ["Intervalo analizado", f"16/10/2026, {hora('baseline_inicio')} a {hora('recuperacion_fin')} (Lima); incidente "
                            f"{hora('incidente_inicio')}–{hora('incidente_fin')} (302 s)."],
    ["Síntoma observable", f"Respuestas lentas e intermitentes: p95 de {ms(B['latencia_p95_s'])} a {n(I['latencia_p95_s'],1,2)} s y "
                           f"{pct(I['ratio_errores_5xx'],1)} de respuestas 503, con tráfico constante y servicio alcanzable (up = 1)."],
    ["Métrica 1 y hallazgo", f"`app_dependency_duration_seconds` (p95): de {ms(B['dependencia_p95_s'])} a {ms(I['dependencia_p95_s'])} "
                             f"(≈ 31×), con ≈ {n(I['dependencia_errores'],1,0)} errores de timeout de inventario-db."],
    ["Métrica 2 y hallazgo", f"`app_http_request_duration_seconds` (p95): de {ms(B['latencia_p95_s'])} a {ms(I['latencia_p95_s'])}; "
                             "por endpoint, la raíz (sin dependencia) se mantuvo en 9,5 ms y checkout (dos llamadas) llegó a 2,49 s."],
    ["Métrica 3 y hallazgo", f"`app_http_requests_total` (ratio 5xx): de {pct(B['ratio_errores_5xx'],3)} a {pct(I['ratio_errores_5xx'],1)}; "
                             "los 503 aparecen solo en /api/productos y /api/checkout."],
    ["Métricas de contraste", f"CPU del host {pct(I['cpu_host_media'],1)} (máx. {pct(I['cpu_host_max_1m'],1)}), CPU de la app "
                              f"invariable, memoria disponible ≥ {pct(I['mem_disponible_min_ratio'],1)}, E/S de disco ≈ 0 %."],
    ["Relación temporal entre señales", "La latencia de la dependencia y la latencia HTTP cambian en el mismo intervalo de 5 s; los "
                                        "errores crecen en rampa durante el minuto siguiente (efecto de la ventana de rate). Las tres "
                                        "señales vuelven a la base al retirar el fallo. CPU y memoria no cambian en ningún momento."],
    ["Hipótesis más consistente", "Fallo de la dependencia inventario-db (latencia alta y timeouts) que se propaga a los endpoints "
                                  "que la consultan."],
    ["Evidencia que contradice alternativas", "CPU: el uso del host bajó (3,8 % → 2,7 %) y el proceso no varió. Memoria: sin descenso. "
                                              "E/S: utilización máxima 0,04 %. Pico de tráfico: la tasa fue 8 req/s en todas las fases. "
                                              "Fallo general del host: la ruta raíz no se degradó."],
    ["Acción inmediata propuesta", "Aislar la dependencia: timeout de 300 ms y cortacircuitos con respuesta degradada (catálogo en "
                                   "caché) y revisar o conmutar inventario-db. No escalar CPU ni memoria."],
    ["Medida preventiva propuesta", "Alertar sobre la latencia y los errores de la dependencia, definir un SLO con alertas por tasa de "
                                    "consumo del presupuesto de errores y añadir trazas distribuidas."],
], [4.4, 12.0], "Hoja de análisis del incidente (formato de la guía)", size=9)

H("7.3. Correlación de señales y contraste de hipótesis", 2)
figure(EV / "incidente/fig_correlacion_senales.png", "Correlación temporal de tres señales del servicio y una señal de contraste",
       "las señales 1 a 3 cambian juntas al inicio y al final del fallo, con la dependencia como primera en reaccionar; la señal 4 "
       "(CPU del host) permanece plana. Hechos observados: coincidencia temporal y selectividad por endpoint. Inferencia: la "
       "dependencia es la causa más consistente.")
table(["Hipótesis de la guía", "Métricas revisadas", "Criterio de contraste", "Resultado"], [
    ["Saturación de CPU", "node_cpu_seconds_total, process_cpu_seconds_total, latencia", "Coincidencia entre presión de CPU y degradación",
     f"**Descartada.** CPU host {pct(I['cpu_host_media'],1)} (base {pct(B['cpu_host_media'],1)}); proceso sin cambios."],
    ["Presión de memoria", "node_memory_MemAvailable_bytes, process_resident_memory_bytes", "Descenso sostenido cerca del incidente",
     f"**Descartada.** Disponible ≥ {pct(I['mem_disponible_min_ratio'],1)}; RSS +1,4 MB."],
    ["Problema de E/S o almacenamiento", "node_disk_io_time_seconds_total", "Cambio anómalo coincidente",
     f"**Descartada.** Utilización máxima {pct(I['disk_io_util_max'])}."],
    ["Fallo de servicio o dependencia", "up, errores, latencia, latencia/errores de inventario-db, concurrencia", "Cambio correlacionado en el servicio",
     "**Confirmada.** up = 1, pero la dependencia aumenta su p95 31× y concentra los errores."],
], [3.2, 4.4, 3.8, 5.0], "Contraste de hipótesis con la evidencia", size=8.5)

H("7.4. Diagnóstico argumentado", 2)
J("Se distinguen tres niveles. **Hechos observados:** el servicio siguió disponible (up = 1) y recibió la misma carga, pero su latencia "
  "y su proporción de errores aumentaron en el mismo intervalo en que lo hizo la latencia de inventario-db; la ruta que no usa la "
  "dependencia no cambió; los recursos del host no variaron. **Inferencia:** la degradación se origina en la dependencia y se "
  "propaga a los endpoints que la consultan; checkout, que hace dos llamadas, sufre más. El aumento de concurrencia es un efecto (ley de "
  "Little), no una causa. **Contexto:** el cambio declarado del experimento fue precisamente la degradación de inventario-db, por lo que el "
  "diagnóstico coincide con la causa real. El marcador `app_fault_injection_active` no se usó como evidencia, para que el razonamiento "
  "dependa solo de las señales que estarían disponibles en un incidente real.")
J("La sospecha inicial de la guía, «presión de recursos», no se confirmó: este caso muestra que una gráfica de latencia por sí sola no "
  "identifica la causa, y que escalar CPU o memoria no habría resuelto el problema.")

H("7.5. Limitaciones del análisis", 2)
for t in [
    "La dependencia es simulada dentro del proceso de la app; en un sistema real habría red, base de datos y conexiones con métricas propias.",
    "Los percentiles se estiman por interpolación lineal dentro de las cubetas del histograma; los valores de p95 son aproximados.",
    "Las ventanas de 1 minuto de rate() suavizan y retrasan las señales hasta unos 60 s; los errores muestran una rampa que no existió en la realidad.",
    "El experimento se ejecutó una vez, en un solo host y durante cinco minutos; no se evaluaron variaciones de carga ni fallos parciales.",
    "node_exporter mide la VM compartida del entorno. Grafana no pudo ejecutarse en ese entorno, por lo que no se incluyen capturas de sus paneles.",
    "increase() extrapola en los bordes de la ventana, por lo que los conteos por fase son aproximados.",
]:
    bullet(t)

# ===========================================================================
# 9 (8). PROPUESTA
# ===========================================================================
H("8. Propuesta de mejora")
H("8.1. Acción correctiva inmediata", 2)
for t in [
    "Aplicar un timeout de 300 ms a cada llamada a inventario-db (≈ 6 veces su p95 normal de 47 ms) para limitar la latencia que se propaga.",
    "Activar un cortacircuitos (*circuit breaker*) que, ante más de un 20 % de fallos en 30 s, responda con datos en caché o una respuesta degradada en lugar de 503.",
    "Revisar inventario-db (conexiones, consultas lentas, despliegues recientes) y, si existe, conmutar a una réplica sana.",
    "No escalar CPU ni memoria de la app: los datos muestran que no son el cuello de botella.",
]:
    bullet(t)
H("8.2. Medidas preventivas", 2)
for t in [
    "Definir un SLO, por ejemplo 99 % de respuestas exitosas y p95 < 300 ms en 30 días, y alertar por tasa de consumo del presupuesto de errores (Beyer et al., 2016).",
    "Reintentos con espera exponencial y aleatoriedad (*jitter*) solo para operaciones idempotentes, con límite de reintentos para no amplificar la carga.",
    "Aislar recursos por dependencia (*bulkhead*): un grupo de hilos dedicado evita que una dependencia lenta agote la capacidad del servicio.",
    "Añadir trazas distribuidas y métricas propias de la base de datos para reducir la incertidumbre en futuros diagnósticos.",
    "Probar periódicamente este escenario (prueba de caos controlada) con los scripts del repositorio.",
]:
    bullet(t)
H("8.3. Reglas de alerta y umbrales justificados", 2)
table(["Alerta", "Condición", "Justificación del umbral"], [
    ["DependenciaLenta", "p95 inventario-db > 250 ms durante 1 m", f"≈ 5× el p95 base ({ms(B['dependencia_p95_s'])}); es la señal más cercana a la causa."],
    ["LatenciaP95Alta", "p95 HTTP > 500 ms durante 1 m", f"≈ 9× el p95 base ({ms(B['latencia_p95_s'])}); mide el impacto en el usuario."],
    ["TasaErroresAlta", "5xx > 5 % durante 1 m", "La base es < 0,2 %; 5 % evita falsos positivos por errores transitorios."],
    ["CPUHostAlta / MemoriaHostBaja", "CPU > 85 % 2 m / memoria < 10 % 2 m", "Detectan presión de recursos; sirven para descartarla rápidamente."],
    ["ConsumoPresupuestoRapido (propuesta)", "ratio 5xx 5 m > 14,4 × 1 % y ratio 1 h > 14,4 × 1 %", "Alerta multiventana por tasa de consumo del SLO del 99 %."],
], [4.0, 5.2, 7.2], "Reglas de alerta implementadas y propuestas")
code("""- alert: ConsumoPresupuestoRapido      # propuesta para el SLO de 99 % de éxito
  expr: |
    ( sum(rate(app_http_requests_total{status=~"5.."}[5m])) / sum(rate(app_http_requests_total[5m])) > 14.4 * 0.01 )
    and
    ( sum(rate(app_http_requests_total{status=~"5.."}[1h])) / sum(rate(app_http_requests_total[1h])) > 14.4 * 0.01 )
  for: 2m
  labels: {severity: critical}""")
H("8.4. Métricas de seguimiento y verificación de la recuperación", 2)
J("Se seguirán `job:app_dependency_latency_seconds:p95_1m`, `job:app_http_latency_seconds:p95_1m`, `job:app_http_errors:ratio_rate1m`, "
  "`app_http_inflight_requests` y la serie `ALERTS`. La recuperación se considera verificada cuando, durante 5 minutos consecutivos: "
  "(1) el ratio de 5xx es < 1 %; (2) el p95 HTTP es < 100 ms; (3) el p95 de la dependencia es < 100 ms; (4) la concurrencia vuelve a "
  "0–2; y (5) no hay alertas en firing. En este laboratorio los cinco criterios se cumplieron en la fase de recuperación (Tabla 6 y Figura 12).")

# ===========================================================================
# 10 (9). CONCLUSIONES
# ===========================================================================
H("9. Conclusiones")
for t in [
    f"La arquitectura Prometheus/TSDB + PromQL + reglas de alerta permitió observar el incidente con resolución de 5 s y trazar cada "
    f"conclusión hasta una consulta concreta: tres alertas de servicio dispararon unos 75 s después del inicio y se resolvieron menos "
    f"de un minuto después de la corrección, sin falsos positivos en las fases normales (Tabla 8).",
    f"La correlación de tres señales identificó la causa: la latencia de inventario-db subió ≈ 31 veces en el mismo intervalo que la "
    f"latencia HTTP (p95 de {ms(B['latencia_p95_s'])} a {n(I['latencia_p95_s'],1,2)} s) y los errores ({pct(I['ratio_errores_5xx'],1)}), "
    f"mientras la ruta sin dependencia no cambió (Tabla 7 y Figura 16).",
    f"La hipótesis de presión de recursos fue descartada con datos: CPU del host de {pct(I['cpu_host_media'],1)}, memoria disponible "
    f"≥ {pct(I['mem_disponible_min_ratio'],1)} y E/S casi nula durante el incidente. Escalar recursos no habría resuelto el problema.",
    "Las métricas de dependencias y los percentiles de latencia son más útiles que las de infraestructura para detectar este tipo de fallo; "
    "por ello se recomienda alertar sobre la dependencia y proteger el servicio con timeouts y cortacircuitos.",
]:
    bullet(t, num=True)

# ===========================================================================
# 11 (10). REFERENCIAS
# ===========================================================================
H("10. Referencias bibliográficas")
refs = [
    "Beyer, B., Jones, C., Petoff, J., & Murphy, N. R. (Eds.). (2016). *Site reliability engineering: How Google runs production systems*. O'Reilly Media. https://sre.google/sre-book/monitoring-distributed-systems/",
    "Docker Inc. (2026). *Docker Compose overview*. Docker Docs. https://docs.docker.com/compose/",
    "GitHub. (2026). *GitHub Codespaces overview*. GitHub Docs. https://docs.github.com/en/codespaces/overview",
    "Grafana Labs. (2026a). *Prometheus data source*. Grafana documentation. https://grafana.com/docs/grafana/latest/datasources/prometheus/",
    "Grafana Labs. (2026b). *Provision Grafana*. Grafana documentation. https://grafana.com/docs/grafana/latest/administration/provisioning/",
    "Majors, C., Fong-Jones, L., & Miranda, G. (2022). *Observability engineering: Achieving production excellence*. O'Reilly Media.",
    "Prometheus Authors. (2026a). *Overview*. Prometheus documentation. https://prometheus.io/docs/introduction/overview/",
    "Prometheus Authors. (2026b). *Data model*. Prometheus documentation. https://prometheus.io/docs/concepts/data_model/",
    "Prometheus Authors. (2026c). *Querying Prometheus*. Prometheus documentation. https://prometheus.io/docs/prometheus/latest/querying/basics/",
    "Prometheus Authors. (2026d). *Storage*. Prometheus documentation. https://prometheus.io/docs/prometheus/latest/storage/",
    "Prometheus Authors. (2026e). *Alerting rules*. Prometheus documentation. https://prometheus.io/docs/prometheus/latest/configuration/alerting_rules/",
    "Prometheus Authors. (2026f). *Histograms and summaries*. Prometheus documentation. https://prometheus.io/docs/practices/histograms/",
    "Prometheus Authors. (2026g). *Node exporter* [Software]. GitHub. https://github.com/prometheus/node_exporter",
]
for r in refs:
    p = doc.add_paragraph()
    add_rich(p, r.replace("*", "**") if False else "")
    import re
    for part in re.split(r"(\*[^*]+\*)", r):
        if part.startswith("*") and part.endswith("*"):
            p.add_run(part[1:-1]).italic = True
        elif part:
            p.add_run(part)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.space_after = Pt(6)

# ===========================================================================
# 12. CIERRE Y METACOGNICIÓN
# ===========================================================================
H("11. Cierre y metacognición")
meta = [
    ("¿Qué métrica fue más útil para detectar o explicar el incidente y por qué?",
     "La latencia p95 de la dependencia (`app_dependency_duration_seconds`). La latencia HTTP y los errores mostraron que el servicio "
     "estaba degradado, pero no por qué. La latencia de inventario-db cambió primero, en el mismo intervalo de 5 s, y en una magnitud "
     "(≈ 31×) que explica por sí sola la latencia de los endpoints que la consultan. Además, su alerta fue la más específica: apunta a "
     "la causa y no solo al síntoma."),
    ("¿Qué hipótesis inicial fue confirmada o descartada por los datos?",
     f"La hipótesis inicial sugerida por el caso, presión de recursos (CPU o memoria), quedó descartada: la CPU del host bajó de "
     f"{pct(B['cpu_host_media'],1)} a {pct(I['cpu_host_media'],1)}, el proceso de la app consumió lo mismo y la memoria disponible no "
     "bajó del 91 %. Se confirmó la hipótesis de fallo de una dependencia, sustentada por la coincidencia temporal y porque la ruta "
     "que no la usa no se degradó."),
    ("¿Qué información adicional habría reducido la incertidumbre del diagnóstico?",
     "Trazas distribuidas que muestren el tiempo de cada petición dentro de inventario-db, métricas propias de la base de datos "
     "(conexiones, consultas lentas, bloqueos) y un registro de cambios o despliegues. También ayudarían histogramas nativos o cubetas "
     "más finas para estimar los percentiles con más precisión y varias repeticiones del experimento para medir la variabilidad."),
    ("¿Cómo automatizaría la detección temprana de un evento similar?",
     "Con las reglas ya probadas (DependenciaLenta, LatenciaP95Alta, TasaErroresAlta) enviadas mediante Alertmanager a un canal del "
     "equipo, más una alerta por tasa de consumo del SLO. Se añadirían sondas sintéticas periódicas contra los endpoints críticos, "
     "un runbook enlazado en cada alerta que indique revisar primero la dependencia y una prueba de caos programada para comprobar "
     "que las alertas siguen funcionando."),
]
for i, (q, a) in enumerate(meta, 1):
    P(f"{i}. {q}", bold=True, after=2)
    J(a)

# ===========================================================================
# 13. AUTOEVALUACIÓN
# ===========================================================================
H("12. Autoevaluación frente a la rúbrica")
table(["Criterio (puntaje)", "Evidencia que sustenta el cumplimiento"], [
    ["Explicación del monitoreo y TSDB (4)", "Sección 3 (Tabla 2, con referencias), Sección 4 (Figura 1 y Tabla 3), configuración de TSDB y recolección."],
    ["Obtención y presentación de métricas (4)", "Inventario (Tabla 5), datos reales exportados con periodos y unidades (Tabla 6), 10 consultas interpretadas (Tabla 7), capturas reales."],
    ["Análisis del fallo (6)", "Línea temporal (Tabla 9), hoja de análisis (Tabla 10), correlación de 3 señales más contraste (Figura 16), 4 hipótesis contrastadas (Tabla 11), diagnóstico y limitaciones."],
    ["Propuesta de respuesta y prevención (3)", "Sección 8: acción inmediata, prevención, reglas con umbrales justificados por la línea base y procedimiento de verificación."],
    ["Comunicación académica y reproducibilidad (3)", "Estructura de la guía, figuras y tablas numeradas con fuente, APA 7, repositorio con scripts y datos."],
], [5.2, 11.2], "Autoevaluación (la calificación corresponde al docente)", size=9)
J("Pendiente declarado: capturas de los paneles de Grafana, que se obtienen al ejecutar `docker compose up` en GitHub Codespaces.")

# ===========================================================================
# ANEXOS
# ===========================================================================
H("13. Anexos")
H("Anexo A. Consultas PromQL completas", 2)
seen = set()
lines = []
for c in CONS:
    if c["fase"] == "incidente" and c["id"] not in seen:
        seen.add(c["id"])
        lines.append(f"# {c['id']} – {c['objetivo']}\n{c['promql'].replace('[302s]', '[Δ]').replace('302s:', 'Δ:')}")
code("\n\n".join(lines), size=7.5)
J("Δ es la duración de cada fase (361 s, 302 s y 360 s). Consultas de los paneles (ventana móvil de 1 minuto):")
code("""sum by (endpoint, status) (rate(app_http_requests_total[1m]))
job:app_http_errors:ratio_rate1m
histogram_quantile(0.95, sum by (le) (rate(app_http_request_duration_seconds_bucket[1m])))
job:app_dependency_latency_seconds:p95_1m
1 - avg(rate(node_cpu_seconds_total{mode="idle"}[1m]))
node_memory_MemAvailable_bytes / node_memory_MemTotal_bytes
ALERTS{alertstate="firing"}""", size=8)

H("Anexo B. Reglas de registro y de alerta (prometheus/alerts.yml)", 2)
code((ROOT / "prometheus/alerts.yml").read_text(encoding="utf-8"), size=7.5)

H("Anexo C. Fragmento de instrumentación y modo de fallo (application/app.py)", 2)
src = (ROOT / "application/app.py").read_text(encoding="utf-8")
a = src.index("def call_dependency")
b = src.index("def instrumented")
code(src[a:b].strip(), size=7.5)

H("Anexo D. Scripts y datos exportados", 2)
table(["Archivo", "Contenido"], [
    ["scripts/validar_entorno.sh", "Verifica herramientas, servicios, objetivos, frescura de muestras, reglas y Grafana"],
    ["scripts/ejecutar_baseline.sh / _incidente.sh / _recuperacion.sh", "Fases del experimento con marcas de tiempo"],
    ["scripts/loadgen.py", "Generador de tráfico controlado (solo contra la app local)"],
    ["scripts/exportar_metricas.py", "Exporta series, consultas por fase y transiciones de alertas"],
    ["scripts/graficar.py / capturas_prometheus.py", "Gráficos con datos reales y capturas de la interfaz"],
    ["resultados/series_experimento.csv", f"{sum(1 for _ in open(RES / 'series_experimento.csv')) - 1} instantes × 21 señales (paso 5 s)"],
    ["resultados/consultas_promql.json", "57 evaluaciones PromQL con objetivo, periodo y resultado"],
    ["resultados/resumen_fases.csv, timeline.csv", "Indicadores por fase y marcas de tiempo"],
    ["resultados/cliente_*.csv", "Vista del cliente: cada petición con estado y latencia"],
    ["grafana/dashboards/catalogo-observabilidad.json", "Dashboard provisionado de 12 paneles"],
], [6.6, 9.8], "Archivos del repositorio", size=8.5)
P(f"Repositorio: {REPO}", bold=True)

H("Anexo E. Evidencias complementarias", 2)
figure(EV / "alertas/prometheus_reglas.png", "Reglas cargadas en Prometheus (estado al finalizar)",
       "las 4 reglas de registro y las 6 de alerta están cargadas y en estado OK, lo que confirma que la configuración se evaluó en el entorno.",
       width=13, source=SRC_UI)
figure(EV / "incidente/prom_p95_incidente.png", "Latencia p95 restringida a la ventana del incidente",
       "durante los 5 minutos del fallo el p95 de 1 min se mantiene entre 2,1 y 2,3 s tras el primer minuto.", width=14, source=SRC_UI)
figure(EV / "recuperacion/prom_p95_recuperacion.png", "Latencia p95 restringida a la ventana de recuperación",
       "el primer minuto muestra la caída desde el valor del incidente y luego el retorno estable al rango de la línea base.",
       width=14, source=SRC_UI)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("Informe:", OUT, "| figuras:", FIG_N[0], "| tablas:", TAB_N[0])
